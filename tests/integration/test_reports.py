from contextlib import closing
from pathlib import Path
from tempfile import TemporaryDirectory
from threading import Event
import unittest
from unittest.mock import patch

from openpyxl import load_workbook
from PySide6.QtCore import QDate, Qt
from PySide6.QtTest import QTest
from PySide6.QtWidgets import QApplication

from attendance_app.application.attendance_service import AttendanceService
from attendance_app.application.export_service import ExportService
from attendance_app.application.people_service import PeopleService
from attendance_app.application.report_service import ReportService
from attendance_app.application.storage_service import StorageService
from attendance_app.configuration.settings import AppConfig
from attendance_app.domain.errors import StorageError
from attendance_app.domain.people import PersonInput
from attendance_app.storage.database import Database
from attendance_app.storage.backup import BackupManager
from attendance_app.storage.repositories.attendance import AttendanceRepository
from attendance_app.storage.repositories.people import PeopleRepository
from attendance_app.ui.formatting import attendance_date
from attendance_app.ui.pages.reports import ReportsPage
from attendance_app.ui.main_window import MainWindow


class ReportsUiTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.owns_app = QApplication.instance() is None
        cls.app = QApplication.instance() or QApplication([])
        cls.app.setLayoutDirection(Qt.LayoutDirection.RightToLeft)

    @classmethod
    def tearDownClass(cls):
        if cls.owns_app:
            import shiboken6
            shiboken6.delete(cls.app)

    def setUp(self):
        self.temp = TemporaryDirectory()
        self.db = Database(Path(self.temp.name) / 'reports.db')
        self.db.initialize()
        self.people = PeopleRepository(self.db)
        self.a = self.people.create(PersonInput('علی', 'کریمی', '0001'))
        self.b = self.people.create(PersonInput('نام', 'دیگر', '0002'))
        self.insert(self.a.id, '2026-10-26')
        self.insert(self.a.id, '2026-10-27')
        self.insert(self.b.id, '2026-10-27')
        self.service = ReportService(AttendanceService(AttendanceRepository(self.db)))
        self.export_service = ExportService(attendance_date)
        self.page = ReportsPage(self.service, self.export_service)
        self.page.resize(944, 640)
        self.page.show()
        self.page.refresh()
        self.settled()

    def tearDown(self):
        self.page.shutdown()
        self.wait(lambda: self.page.presenter.worker is None and self.page.presenter.export_worker is None)
        self.page.close()
        self.page.deleteLater()
        self.app.processEvents()
        self.temp.cleanup()

    def insert(self, person_id, day):
        with self.db.connect(write=True) as connection:
            connection.execute('''INSERT INTO attendance
                (person_id,attendance_date,attendance_time,observed_at,created_at)
                VALUES (?,?,'08:30:00+03:30',?,?)''',
                (person_id, day, day + 'T05:00:00+00:00', day + 'T05:00:01+00:00'))

    def wait(self, predicate):
        for _ in range(300):
            self.app.processEvents()
            if predicate():
                return
            QTest.qWait(20)
        self.fail('Reports did not settle')

    def settled(self):
        self.wait(lambda: self.page.presenter.worker is None)

    def export(self, path):
        with patch('attendance_app.ui.pages.reports.QFileDialog.getSaveFileName', return_value=(str(path), '')):
            self.page.export_button.click()
        self.wait(lambda: self.page.presenter.export_worker is None)

    def test_filters_summaries_reset_dates_and_empty(self):
        self.assertEqual([self.page.status.itemText(i) for i in range(self.page.status.count())],
                         ['همه وضعیتها', 'حاضر'])
        self.assertEqual((self.page.snapshot.record_count, self.page.snapshot.people_count), (3, 2))
        self.people.archive(self.b.id)
        self.page.refresh()
        self.settled()
        self.assertIn('بایگانی‌شده', self.page.person.itemText(self.page.person.findData(self.b.id)))
        self.page.person.setCurrentIndex(self.page.person.findData(self.a.id))
        self.page.date_mode.setCurrentIndex(2)
        self.page.start_date.setDate(QDate(2026, 10, 27))
        self.page.end_date.setDate(QDate(2026, 10, 27))
        self.page.status.setCurrentIndex(1)
        self.settled()
        self.assertEqual(self.page.snapshot.record_count, 1)
        self.assertEqual(self.page.model.data(self.page.model.index(0, 2)), '1405/08/05')
        self.page.start_date.setDate(QDate(2026, 10, 28))
        self.assertIsNone(self.page.snapshot)
        self.assertFalse(self.page.export_button.isEnabled())
        self.page.reset_button.click()
        self.settled()
        self.assertEqual(self.page.snapshot.record_count, 3)
        self.assertEqual(self.page.status.currentIndex(), 0)
        self.page.search.setText('missing')
        self.wait(lambda: self.page.snapshot is not None)
        self.assertEqual(self.page.snapshot.record_count, 0)
        self.assertFalse(self.page.export_button.isEnabled())

    def test_sorted_snapshot_export_ignores_later_database_changes(self):
        self.page.table.sortByColumn(1, Qt.SortOrder.AscendingOrder)
        expected = tuple(self.page.model.records)
        self.people.update(self.a.id, PersonInput('جدید', 'نام', '9999'))
        self.insert(self.b.id, '2026-10-28')
        path = Path(self.temp.name) / 'گزارش.xlsx'
        self.export(path)
        with closing(load_workbook(path)) as book:
            self.assertEqual(book.worksheets[1].max_row, 4)
            self.assertEqual([book.worksheets[1].cell(i, 2).value for i in range(2, 5)],
                             [r.public_person_id for r in expected])
        self.page.refresh()
        self.settled()
        self.assertEqual(self.page.snapshot.record_count, 4)
        self.assertIn('9999', [r.public_person_id for r in self.page.model.records])

    def test_save_cancel_declined_replace_and_error_recovery(self):
        with patch('attendance_app.ui.pages.reports.QFileDialog.getSaveFileName', return_value=('', '')):
            self.page.export_button.click()
        self.assertIsNone(self.page.presenter.export_worker)
        path = Path(self.temp.name) / 'report.xlsx'
        path.write_bytes(b'old')
        with patch('attendance_app.ui.pages.reports.confirm', return_value=False):
            self.export(path)
        self.assertEqual(path.read_bytes(), b'old')
        with patch('attendance_app.ui.pages.reports.confirm', return_value=True), \
                patch.object(self.export_service.writer, 'write', side_effect=PermissionError()):
            self.export(path)
        self.assertIn('خطا', self.page.export_feedback.text())
        self.assertEqual(path.read_bytes(), b'old')
        with patch('attendance_app.ui.pages.reports.confirm', return_value=True):
            self.export(path)
        self.assertIn('ذخیره شد', self.page.export_feedback.text())

    def test_read_failure_retry_and_slow_latest_request(self):
        with patch.object(self.service, 'read', side_effect=StorageError()):
            self.page.refresh()
            self.settled()
        self.assertIsNone(self.page.snapshot)
        self.assertIn('—', self.page.summary.text())
        self.assertFalse(self.page.export_button.isEnabled())
        entered, release = Event(), Event()
        original = self.service.read
        calls = []
        def slow(filters):
            calls.append(filters)
            result = original(filters)
            if len(calls) == 1:
                entered.set()
                if not release.wait(5):
                    raise RuntimeError('timeout')
            return result
        try:
            with patch.object(self.service, 'read', side_effect=slow):
                self.page.refresh()
                self.wait(entered.is_set)
                self.page.search.setText('0002')
                self.page.refresh()
                self.assertIsNone(self.page.snapshot)
                release.set()
                self.settled()
            self.assertEqual(len(calls), 2)
            self.assertEqual(self.page.snapshot.record_count, 1)
            self.assertEqual(self.page.snapshot.filters.query, '0002')
        finally:
            release.set()

    def test_export_busy_and_shutdown_waits(self):
        entered, release = Event(), Event()
        original = self.export_service.export
        def slow(*args, **kwargs):
            entered.set()
            if not release.wait(5):
                raise RuntimeError('timeout')
            return original(*args, **kwargs)
        try:
            with patch.object(self.export_service, 'export', side_effect=slow):
                self.page.presenter.export(self.page.snapshot, Path(self.temp.name) / 'report.xlsx')
                self.page.update_actions()
                self.wait(entered.is_set)
                self.assertFalse(self.page.export_button.isEnabled())
                self.assertFalse(self.page.presenter.export(self.page.snapshot, 'unused.xlsx'))
                self.assertFalse(self.page.shutdown())
                release.set()
                self.wait(lambda: self.page.presenter.export_worker is None)
                self.assertTrue(self.page.shutdown())
        finally:
            release.set()

    def test_file_dialog_nested_refresh_cannot_export_obsolete_snapshot(self):
        path = Path(self.temp.name) / 'obsolete.xlsx'
        def choose(*args, **kwargs):
            self.page.refresh()
            return str(path), ''
        with patch('attendance_app.ui.pages.reports.QFileDialog.getSaveFileName', side_effect=choose):
            self.page.export_button.click()
        self.settled()
        self.assertFalse(path.exists())
        self.assertIsNone(self.page.presenter.export_worker)

    def test_navigation_restore_and_main_close_waits_for_export(self):
        backup = BackupManager(self.db)
        backup_path = Path(self.temp.name) / 'backup.db'
        backup.backup(backup_path)
        window = MainWindow(AppConfig(), PeopleService(self.people), StorageService(backup),
                            attendance_service=self.service.attendance,
                            report_service=self.service, export_service=self.export_service)
        entered, release = Event(), Event()
        original = self.export_service.export
        def slow(*args, **kwargs):
            entered.set()
            if not release.wait(5):
                raise RuntimeError('timeout')
            return original(*args, **kwargs)
        try:
            window.show()
            window.navigate_to('reports')
            self.wait(lambda: window.reports_page.snapshot is not None)
            self.assertIs(window.pages.currentWidget(), window.reports_page)
            self.insert(self.b.id, '2026-10-28')
            window.navigate_to('dashboard')
            window.navigate_to('reports')
            self.wait(lambda: window.reports_page.snapshot is not None)
            self.assertEqual(window.reports_page.snapshot.record_count, 4)
            backup.restore(backup_path, confirmed=True)
            window.pages.widget(5).restored.emit()
            self.assertIsNone(window.reports_page.snapshot)
            self.wait(lambda: window.reports_page.snapshot is not None)
            self.assertEqual(window.reports_page.snapshot.record_count, 3)
            with patch.object(self.export_service, 'export', side_effect=slow):
                window.reports_page.presenter.export(window.reports_page.snapshot,
                                                     Path(self.temp.name) / 'closing.xlsx')
                self.wait(entered.is_set)
                window.close()
                self.assertTrue(window.isVisible())
                release.set()
                self.wait(lambda: not window.isVisible())
            self.assertTrue((Path(self.temp.name) / 'closing.xlsx').exists())
        finally:
            release.set()
            window.close()
            self.wait(lambda: all(page.presenter.worker is None for page in
                (window.history_page, window.dashboard_page, window.reports_page))
                and window.reports_page.presenter.export_worker is None)
            window.close()
            window.deleteLater()
            self.app.processEvents()
