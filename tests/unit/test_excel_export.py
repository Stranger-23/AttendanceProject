from dataclasses import replace
from contextlib import closing
from datetime import datetime, timezone
from pathlib import Path
from tempfile import TemporaryDirectory
import unittest
from unittest.mock import patch

from openpyxl import load_workbook

from attendance_app.application.export_service import ExportService
from attendance_app.domain.attendance import AttendanceHistoryRecord
from attendance_app.domain.errors import ExportError
from attendance_app.domain.reports import ReportFilters, ReportSnapshot
from attendance_app.reporting.excel import ExcelWriter, HEADERS
from attendance_app.ui.formatting import attendance_date


class ExcelExportTests(unittest.TestCase):
    def setUp(self):
        self.temp = TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.path = Path(self.temp.name) / 'گزارش.xlsx'
        self.rows = tuple(AttendanceHistoryRecord(i, i, identifier, name, False,
                           '2025-03-20', '23:59:59-04:00', 'present')
                          for i, (identifier, name) in enumerate((('00001', '=1+1'),
                          ('12345678901234567890', 'https://example.org'), ('@name', 'علی کریمی')), 1))
        self.snapshot = ReportSnapshot(ReportFilters(query='=test'), self.rows,
                           datetime(2026, 10, 27, 8, tzinfo=timezone.utc), 'همه افراد')
        self.service = ExportService(attendance_date)

    def test_workbook_values_types_format_and_exact_order(self):
        self.service.export(self.snapshot, self.path)
        with closing(load_workbook(self.path)) as book:
            self.assertEqual(book.sheetnames, ['اطلاعات گزارش', 'حضورهای ثبت‌شده'])
            sheet = book.worksheets[1]
            self.assertEqual(tuple(c.value for c in sheet[1]), HEADERS)
            self.assertTrue(all(s.sheet_view.rightToLeft for s in book))
            self.assertEqual(sheet.freeze_panes, 'A2')
            self.assertEqual(sheet.auto_filter.ref, 'A1:E4')
            self.assertEqual(sheet.max_column, 5)
            for index, record in enumerate(self.rows, 2):
                self.assertEqual(tuple(c.value for c in sheet[index]),
                    (record.full_name, record.public_person_id, '1403/12/30', '23:59:59', 'حاضر'))
                self.assertTrue(all(c.data_type == 's' and c.hyperlink is None for c in sheet[index]))
            self.assertEqual(book.worksheets[0]['B3'].value, '=test')
            self.assertEqual(book.worksheets[0]['B3'].data_type, 's')
            self.assertEqual(book.worksheets[0]['B9'].value, 3)
            self.assertEqual(book.worksheets[0]['B9'].data_type, 'n')
        self.assertEqual(list(self.path.parent.glob('.attendance-*')), [])

    def test_refused_overwrite_and_appearing_destination_are_preserved(self):
        self.path.write_bytes(b'old')
        with self.assertRaises(ExportError):
            self.service.export(self.snapshot, self.path)
        self.assertEqual(self.path.read_bytes(), b'old')
        self.path.unlink()
        original = self.service.writer.write
        def race(*args):
            original(*args)
            self.path.write_bytes(b'appeared')
        with patch.object(self.service.writer, 'write', side_effect=race):
            with self.assertRaises(ExportError):
                self.service.export(self.snapshot, self.path)
        self.assertEqual(self.path.read_bytes(), b'appeared')
        self.assertEqual(list(self.path.parent.glob('.attendance-*')), [])

    def test_generation_and_publication_failure_preserve_old_file(self):
        self.path.write_bytes(b'old')
        for target in ('writer', 'replace'):
            context = patch.object(self.service.writer, 'write', side_effect=OSError()) if target == 'writer' else \
                      patch('attendance_app.application.export_service.os.replace', side_effect=PermissionError())
            with context, self.assertRaises(ExportError):
                self.service.export(self.snapshot, self.path, overwrite=True)
            self.assertEqual(self.path.read_bytes(), b'old')
            self.assertEqual(list(self.path.parent.glob('.attendance-*')), [])
        self.service.export(self.snapshot, self.path, overwrite=True)
        self.assertTrue(self.path.read_bytes().startswith(b'PK'))

    @unittest.skipUnless(__import__('os').name == 'nt', 'Windows file sharing semantics')
    def test_open_destination_is_preserved_on_windows(self):
        self.path.write_bytes(b'old')
        with self.path.open('rb'):
            with self.assertRaises(ExportError):
                self.service.export(self.snapshot, self.path, overwrite=True)
        self.assertEqual(self.path.read_bytes(), b'old')
        self.assertEqual(list(self.path.parent.glob('.attendance-*')), [])

    def test_empty_limits_and_no_silent_truncation(self):
        with self.assertRaises(ExportError):
            self.service.export(replace(self.snapshot, records=()), self.path)
        with patch('attendance_app.reporting.excel.MAX_RECORDS', 2), self.assertRaises(ExportError):
            self.service.export(self.snapshot, self.path)
        too_long = replace(self.rows[0], full_name='x' * 32768)
        with self.assertRaises(ExportError):
            self.service.export(replace(self.snapshot, records=(too_long,)), self.path)
        many = tuple(replace(self.rows[0], id=i, public_person_id=f'{i:06}') for i in range(5003))
        self.service.export(replace(self.snapshot, records=many), self.path)
        with closing(load_workbook(self.path, read_only=True)) as book:
            self.assertEqual(book.worksheets[1].max_row, 5004)
            self.assertEqual(book.worksheets[1]['B5004'].value, '005002')
